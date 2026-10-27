"""
Renders the generated process documentation (a small markdown subset: headings,
lists, tables, code blocks, **bold** and `code`) as a Word or PDF file.
"""
import io
import re
from typing import Iterator, List, Tuple

from docx import Document
from docx.shared import Pt, RGBColor
from fpdf import FPDF

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")

# Core PDF fonts are Latin-1 only
PDF_CHARS = {"→": "->", "—": "-", "–": "-", "•": "-", "…": "...", "‘": "'", "’": "'", "“": '"', "”": '"', "≥": ">=", "≤": "<="}


def _blocks(md: str) -> Iterator[Tuple[str, object]]:
    """Yields (kind, data): h1/h2/p -> text, li -> (level, marker, text), table -> rows, code -> text."""
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code: List[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            yield "code", "\n".join(code)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-|\s]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i][1:-1].split("|")])
                i += 1
            yield "table", rows
            continue
        elif line.startswith("# "):
            yield "h1", line[2:]
        elif line.startswith("## "):
            yield "h2", line[3:]
        elif re.match(r"^\s*(\d+\.|-)\s", line):
            m = re.match(r"^(\s*)(\d+\.|-)\s(.*)$", line)
            yield "li", (1 if len(m.group(1)) >= 2 else 0, m.group(2), m.group(3))
        elif line.strip():
            yield "p", line
        i += 1


def _runs(text: str) -> Iterator[Tuple[str, str]]:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            yield "bold", part[2:-2]
        elif part.startswith("`"):
            yield "code", part[1:-1]
        else:
            yield "plain", part


def _strip(text: str) -> str:
    return "".join(t for _, t in _runs(text))


# --------------------------------------------------------------------------- Word

def _add_runs(par, text: str) -> None:
    for kind, t in _runs(text):
        run = par.add_run(t)
        if kind == "bold":
            run.bold = True
        elif kind == "code":
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def to_docx(md: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for kind, data in _blocks(md):
        if kind == "h1":
            doc.add_heading(data, level=0)
        elif kind == "h2":
            doc.add_heading(data, level=1)
        elif kind == "p":
            _add_runs(doc.add_paragraph(), data)
        elif kind == "li":
            level, marker, text = data
            style = "List Number" if marker[0].isdigit() else ("List Bullet 2" if level else "List Bullet")
            _add_runs(doc.add_paragraph(style=style), text)
        elif kind == "table" and data:
            table = doc.add_table(rows=0, cols=len(data[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(data):
                cells = table.add_row().cells
                for c, value in enumerate(row[: len(cells)]):
                    par = cells[c].paragraphs[0]
                    if r == 0:
                        par.add_run(_strip(value)).bold = True
                    else:
                        _add_runs(par, value)
            doc.add_paragraph()
        elif kind == "code":
            par = doc.add_paragraph()
            run = par.add_run(data)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF

def _latin1(text: str) -> str:
    for k, v in PDF_CHARS.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def _pdf_md(text: str) -> str:
    # fpdf2's markdown mode understands **bold**; render `code` as plain text
    return _latin1(text.replace("`", ""))


def to_pdf(md: str, title: str) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_title(_latin1(title))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    width = pdf.w - pdf.l_margin - pdf.r_margin
    for kind, data in _blocks(md):
        pdf.set_x(pdf.l_margin)
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 18)
            pdf.set_text_color(15, 50, 90)
            pdf.multi_cell(width, 9, _latin1(data), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif kind == "h2":
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 13)
            pdf.set_text_color(0, 102, 170)
            pdf.multi_cell(width, 7, _latin1(data), new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(200, 210, 220)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + width, pdf.get_y())
            pdf.ln(2)
        elif kind == "p":
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(30, 30, 30)
            pdf.multi_cell(width, 5.5, _pdf_md(data), markdown=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "li":
            level, marker, text = data
            indent = 5 + level * 7
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(30, 30, 30)
            pdf.set_x(pdf.l_margin + indent)
            bullet = marker if marker[0].isdigit() else "-"
            pdf.multi_cell(width - indent, 5.5, _pdf_md(f"{bullet} {text}"), markdown=True, new_x="LMARGIN", new_y="NEXT")
        elif kind == "table" and data:
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(30, 30, 30)
            with pdf.table(width=width, first_row_as_headings=True, line_height=5.5) as table:
                for row in data:
                    r = table.row()
                    for value in row:
                        r.cell(_latin1(_strip(value)))
            pdf.ln(2)
        elif kind == "code":
            pdf.set_font("Courier", "", 8)
            pdf.set_text_color(60, 60, 60)
            pdf.set_fill_color(242, 245, 248)
            pdf.multi_cell(width, 4.2, _latin1(data), fill=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
    return bytes(pdf.output())
